import io
from pathlib import Path

import docx
import pdfplumber

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"
MAX_BYTES = 10 * 1024 * 1024

MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class UnsupportedResume(ValueError):
    pass


def validate_upload(filename: str, data: bytes) -> str:
    """Returns the lowercase extension (.pdf/.docx) or raises UnsupportedResume."""
    ext = Path(filename).suffix.lower()
    if ext not in MIME:
        raise UnsupportedResume("Only PDF and DOCX resumes are supported")
    if not data:
        raise UnsupportedResume("File is empty")
    if len(data) > MAX_BYTES:
        raise UnsupportedResume("File is larger than 10 MB")
    magic = PDF_MAGIC if ext == ".pdf" else ZIP_MAGIC
    if not data.startswith(magic):
        raise UnsupportedResume("File content does not match its extension")
    return ext


def extract_text(data: bytes, ext: str) -> str:
    if ext == ".pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(p.extract_text() or "") for p in pdf.pages]
        return "\n\n".join(pages).strip()
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()
